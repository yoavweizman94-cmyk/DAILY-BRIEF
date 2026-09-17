# -*- coding: utf-8 -*-
"""מחירי סחורות, עקומי חוזים, מחירי ייחוס חודשיים וכותרות — לעמוד commodities.html.

ההגדרות ב-config/commodities.yaml (שם גם ההסבר למה ארבעה מקורות):
  · Trading Economics /markets/commodities — מחיר ושינויים, 109 סחורות.
  · yfinance — חוזים עתידיים לפי חודש אספקה, לעקום.
  · הבנק העולמי, Pink Sheet חודשי — מחירי ייחוס לפי אזור, עם התיאור הרשמי.
  · Google News ו-RSS — היצע, ביקוש, מאקרו ועסקאות.

**כל מקור נכשל לבד.** מנוי שלא ענה, חוזה שלא נסחר או גיליון שהוחלף — העמוד
נבנה ממה שכן הגיע, ו-state.json אומר מה חסר.

פלט (output/commodities/):
  prices/<YYYY-MM-DD>.json   — תמונת המחירים והעקומים של היום (נדרסת במהלך היום)
  worldbank.json             — הסדרות החודשיות ותיאורן הרשמי
  items/<YYYY-MM-DD>.jsonl   — כותרות {id, ts, title, source, url, region, themes, companies, snippet}
  state.json                 — מצב הריצה

GITHUB_OUTPUT: new=<מספר>, analyze=yes|no, changed=yes|no
"""
from __future__ import annotations

import hashlib
import io
import json
import os
import re
import sys
import time
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from _tls import harden  # noqa: E402
harden()

import yaml  # noqa: E402
from curl_cffi import requests as creq  # noqa: E402

# אותם כלים כמו בגיליון הרכב: קריאת פידים, Google News, התאמת מילים ודה-דופ.
from auto_pull import gn_url, hits, matchers, norm, parse_feed, title_key  # noqa: E402
from _env import load_env  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
CFG = ROOT / "config" / "commodities.yaml"
COMPANIES = ROOT / "config" / "companies.yaml"
OUT = ROOT / "output" / "commodities"
TE = "https://api.tradingeconomics.com/markets/commodities"

LOOKBACK_H = 72
FIRST_RUN_H = 24 * 5
KEEP_DAYS = 8
MIN_NEW = int(os.environ.get("COMMOD_MIN_NEW", "5"))
MIN_HOURS = float(os.environ.get("COMMOD_MIN_HOURS", "10"))
MONTH_CODES = "FGHJKMNQUVXZ"
TE_FIELDS = ("Name", "Group", "Last", "unit", "Date", "LastUpdate", "DailyPercentualChange",
             "WeeklyPercentualChange", "MonthlyPercentualChange", "YearlyPercentualChange",
             "YTDPercentualChange", "yesterday", "lastWeek", "lastMonth", "lastYear", "startYear",
             "allTimeHigh", "allTimeHighDate", "allTimeLow", "allTimeLowDate", "decimals")


def _tz():
    try:
        from zoneinfo import ZoneInfo
        return ZoneInfo("Asia/Jerusalem")
    except Exception:  # noqa: BLE001 — Windows בלי tzdata
        return timezone(timedelta(hours=3))


IL = _tz()


# --------------------------------------------------------------------------
# Trading Economics

def te_snapshot() -> tuple[dict, str | None]:
    key = os.environ.get("TRADINGECONOMICS_KEY")
    if not key:
        return {}, "TRADINGECONOMICS_KEY חסר"
    import requests
    for attempt in range(3):
        try:
            r = requests.get(TE, params={"c": key, "f": "json"}, timeout=60)
        except Exception as e:  # noqa: BLE001
            err = f"{type(e).__name__}"
            time.sleep(6)
            continue
        if r.status_code == 200 and (r.headers.get("content-type") or "").startswith("application/json"):
            rows = r.json()
            return {x["Symbol"]: {k: x.get(k) for k in TE_FIELDS}
                    for x in rows if isinstance(x, dict) and x.get("Symbol")}, None
        # TE מחזיר 403 גם על חנק קצב; "slow down" בגוף — ממתינים ומנסים שוב.
        err = f"HTTP {r.status_code}: {r.text[:80]}"
        time.sleep(8 if "slow down" in r.text.lower() else 3)
    return {}, err


# --------------------------------------------------------------------------
# עקומי חוזים

def contract_for(root: str, suffix: str, allowed: str, today: date, ahead: int) -> tuple[str, str]:
    """החוזה הנסחר הראשון שחודש האספקה שלו לפחות `ahead` חודשים קדימה."""
    base = today.year * 12 + today.month - 1 + ahead
    for k in range(36):
        y, m = divmod(base + k, 12)
        if MONTH_CODES[m] in allowed:
            return f"{root}{MONTH_CODES[m]}{y % 100:02d}.{suffix}", f"{y:04d}-{m + 1:02d}"
    raise ValueError(f"אין חודש חוזה ל-{root}")


def curves(cfg: dict, today: date) -> tuple[dict, str | None]:
    try:
        import yfinance as yf
    except ImportError:
        return {}, "yfinance לא מותקן"
    meta, tickers = {}, []
    for c in cfg.get("curves") or []:
        front = f"{c['root']}=F"
        meta[front] = (c, "front", None)
        tickers.append(front)
        for h in cfg.get("curve_horizons") or [3, 6, 12, 24]:
            t, month = contract_for(c["root"], c["suffix"], c["months"], today, int(h))
            if t not in meta:
                meta[t] = (c, int(h), month)
                tickers.append(t)
    try:
        data = yf.download(tickers, period="10d", interval="1d", progress=False, group_by="ticker",
                           threads=True, auto_adjust=False)
    except Exception as e:  # noqa: BLE001
        return {}, f"yfinance: {type(e).__name__}"
    out: dict = {}
    for t in tickers:
        try:
            col = data[t]["Close"].dropna()
        except Exception:  # noqa: BLE001 — חוזה שלא הוחזר
            continue
        if col.empty:
            continue
        c, h, month = meta[t]
        e = out.setdefault(c["root"], {"label": c["label"], "unit": c["unit"], "points": []})
        point = {"price": round(float(col.iloc[-1]), 4), "date": str(col.index[-1].date())}
        if h == "front":
            e["front"] = point
        else:
            e["points"].append({"ahead": h, "contract": t.split(".")[0], "month": month, **point})
    for e in out.values():
        e["points"].sort(key=lambda p: p["ahead"])
        f = (e.get("front") or {}).get("price")
        for p in e["points"]:
            p["vs_front_pct"] = round((p["price"] / f - 1) * 100, 2) if f else None
    return {k: v for k, v in out.items() if v.get("front")}, None


# --------------------------------------------------------------------------
# הבנק העולמי

def worldbank(cfg: dict, s, previous: dict) -> tuple[dict, str | None]:
    wcfg = cfg.get("worldbank") or {}
    try:
        page = s.get(wcfg.get("landing"), timeout=60).text
        links = re.findall(r'https?://[^"\']+CMO-Historical-Data-Monthly\.xlsx', page)
    except Exception as e:  # noqa: BLE001
        return previous, f"דף הבנק העולמי: {type(e).__name__}"
    if not links:
        return previous, "לא נמצא קישור לגיליון החודשי"
    url = links[0]
    fetched = previous.get("fetched_at") or ""
    # אותו גיליון, ונקרא בשלושת הימים האחרונים — אין מה למשוך שוב.
    if previous.get("url") == url and fetched[:10] >= (date.today() - timedelta(days=3)).isoformat():
        return previous, None
    try:
        import openpyxl
        x = s.get(url, timeout=120)
        wb = openpyxl.load_workbook(io.BytesIO(x.content), read_only=True, data_only=True)
        rows = list(wb["Monthly Prices"].iter_rows(values_only=True))
        desc_rows = [[str(c) for c in r if c is not None] for r in wb["Description"].iter_rows(values_only=True)]
    except Exception as e:  # noqa: BLE001
        return previous, f"גיליון הבנק העולמי: {type(e).__name__}"
    updated = next((str(r[0]) for r in rows[:6] if r and str(r[0] or "").startswith("Updated")), "")
    head_i = next(i for i, r in enumerate(rows) if r and any(str(c or "").startswith("Crude oil") for c in r))
    header = [str(c).strip() if c is not None else "" for c in rows[head_i]]
    units = [str(c).strip() if c is not None else "" for c in rows[head_i + 1]]
    data = [r for r in rows if r and re.fullmatch(r"\d{4}M\d{2}", str(r[0] or ""))]
    descs = [" ".join(t for t in r if t.strip("* ")) for r in desc_rows]
    n = int(wcfg.get("months") or 36)
    out = {}
    for sdef in wcfg.get("series") or []:
        col = sdef["col"].strip()
        if col not in header:
            continue
        i = header.index(col)
        pts = []
        for r in data[-n:]:
            v = r[i]
            if isinstance(v, (int, float)):
                pts.append({"month": f"{str(r[0])[:4]}-{str(r[0])[5:7]}", "value": round(float(v), 4)})
        # **התחלה מדויקת ולא התאמה עמומה.** התאמה לפי מילים שייכה ל-DAP את תיאור
        # שמן הדקלים (שבו כתוב "Crude, DAP") ולאוריאה תיאור של שרימפס.
        prefix = sdef.get("desc") or ""
        desc = next((d for d in descs if prefix and d.startswith(prefix)), "")
        out[sdef["key"]] = {"col": col, "label": sdef.get("label"), "region": sdef.get("region"),
                            "group": sdef.get("group"), "unit": units[i].strip("()"), "points": pts,
                            "desc": desc[:400]}
    return {"url": url, "updated": updated, "fetched_at": datetime.now(IL).isoformat(timespec="minutes"),
            "series": out}, None


# --------------------------------------------------------------------------
# כותרות

def read_items(days: int = KEEP_DAYS) -> list[dict]:
    rows = []
    d = OUT / "items"
    if d.exists():
        for p in sorted(d.glob("*.jsonl"), reverse=True)[:days]:
            for line in p.read_text(encoding="utf-8").splitlines():
                if line.strip():
                    try:
                        rows.append(json.loads(line))
                    except ValueError:
                        pass
    return rows


def write_items(new: list[dict]) -> None:
    by_day: dict[str, list[dict]] = {}
    for r in new:
        day = datetime.fromisoformat(r["ts"]).astimezone(IL).date().isoformat()
        by_day.setdefault(day, []).append(r)
    (OUT / "items").mkdir(parents=True, exist_ok=True)
    for day, rows in by_day.items():
        p = OUT / "items" / f"{day}.jsonl"
        old = [json.loads(x) for x in p.read_text(encoding="utf-8").splitlines() if x.strip()] if p.exists() else []
        seen = {o["id"] for o in old}
        merged = old + [r for r in rows if r["id"] not in seen]
        merged.sort(key=lambda r: r["ts"], reverse=True)
        with p.open("w", encoding="utf-8", newline="\n") as f:
            for r in merged:
                f.write(json.dumps(r, ensure_ascii=False) + "\n")


def headlines(cfg: dict, known: set[str], now: datetime) -> tuple[list[dict], list[str], dict]:
    themes = [(t["slug"], matchers(t.get("keywords"))) for t in cfg.get("themes") or []]
    business = matchers(cfg.get("business_words"))
    ex_sources = [norm(x) for x in cfg.get("exclude_sources") or []]
    ex_titles = [norm(x) for x in cfg.get("exclude_title") or []]
    world_ok = [norm(x) for x in cfg.get("world_sources") or []]
    companies, bad = [], []
    for c in cfg.get("mentions") or []:
        if c["name"] not in known:
            bad.append(c["name"])
        companies.append((c["name"], matchers([c["name"], *(c.get("aliases") or [])], whole_word_hebrew=True)))
    if bad:
        print("::warning title=חברות בעמוד הסחורות שאינן ברשימת הכיסוי::" + ", ".join(bad))

    sources = [(f["name"], f["url"], f.get("region") or "world", False) for f in cfg.get("feeds") or []]
    for region, queries in (cfg.get("google_news") or {}).items():
        for q in queries or []:
            sources.append((f"Google News: {q}", gn_url(q, region), region, True))

    existing = read_items()
    cutoff = now - timedelta(hours=LOOKBACK_H if existing else FIRST_RUN_H)
    seen = {r["id"] for r in existing}
    session = creq.Session(impersonate="chrome")
    fresh: dict[str, dict] = {}
    failures, dropped = [], {"ישן": 0, "מקור": 0, "לא שוק": 0, "בלי נושא": 0}
    for name, url, region, gn in sources:
        try:
            r = session.get(url, timeout=30, allow_redirects=True)
            r.raise_for_status()
            rows = parse_feed(r.text, name, gn)
        except Exception as e:  # noqa: BLE001
            failures.append(f"{name} ({type(e).__name__})")
            continue
        finally:
            time.sleep(1.2 if gn else 0.3)
        for x in rows:
            ts = x["_ts"] or now
            if ts < cutoff:
                dropped["ישן"] += 1
                continue
            src = norm(x["source"])
            if any(e and e in src for e in ex_sources) or \
                    (gn and region == "world" and not any(w and w in src for w in world_ok)):
                dropped["מקור"] += 1
                continue
            if any(e and e in norm(x["title"]) for e in ex_titles):
                dropped["מקור"] += 1
                continue
            blob = norm(f"{x['title']} {x['snippet']}")
            if region == "world" and not hits(business, blob):
                dropped["לא שוק"] += 1
                continue
            th = [slug for slug, pats in themes if hits(pats, blob)]
            # "עסקאות" ו"מאקרו" הם חתך, לא נושא: כותרת שיש בה רק אותם אינה על סחורה.
            if not [t for t in th if t not in ("deals", "macro")]:
                dropped["בלי נושא"] += 1
                continue
            rid = hashlib.sha1(title_key(x["title"]).encode("utf-8")).hexdigest()[:16]
            if rid in seen:
                continue
            item = {"id": rid, "ts": ts.astimezone(timezone.utc).isoformat(timespec="seconds"),
                    "title": x["title"], "source": x["source"], "url": x["url"], "region": region,
                    "themes": th, "companies": [n for n, pats in companies if hits(pats, blob)],
                    "snippet": x["snippet"]}
            prev = fresh.get(rid)
            if prev:
                prev["themes"] = sorted(set(prev["themes"]) | set(th))
                prev["companies"] = sorted(set(prev["companies"]) | set(item["companies"]))
                if region == "il":
                    prev["region"] = "il"
                continue
            fresh[rid] = item
    return list(fresh.values()), failures, {"sources": len(sources), "dropped": dropped, "existing": existing}


# --------------------------------------------------------------------------

def last_analysis() -> dict | None:
    d = OUT / "analyses"
    if not d.exists():
        return None
    for p in sorted(d.glob("*.jsonl"), reverse=True):
        for line in reversed([x for x in p.read_text(encoding="utf-8").splitlines() if x.strip()]):
            try:
                return json.loads(line)
            except ValueError:
                continue
    return None


def main() -> int:
    load_env()
    cfg = yaml.safe_load(CFG.read_text(encoding="utf-8")) or {}
    known = {(c.get("name_he") or "").strip()
             for c in (yaml.safe_load(COMPANIES.read_text(encoding="utf-8")) or {}).get("companies") or []}
    missing = sorted({n for e in cfg.get("exposures") or [] for g in e.get("groups") or []
                      for n in g.get("companies") or [] if n not in known})
    if missing:
        print("::warning title=חשיפות לחברות שאינן ברשימת הכיסוי::" + ", ".join(missing))

    now = datetime.now(timezone.utc)
    today = datetime.now(IL).date()
    OUT.mkdir(parents=True, exist_ok=True)
    failures = []

    te, err = te_snapshot()
    if err:
        failures.append(f"Trading Economics: {err}")
    cv, err = curves(cfg, today)
    if err:
        failures.append(err)
    snap_p = OUT / "prices" / f"{today.isoformat()}.json"
    prices_changed = False
    if te or cv:
        old = {}
        if snap_p.exists():
            try:
                old = json.loads(snap_p.read_text(encoding="utf-8"))
            except ValueError:
                old = {}
        snap = {"date": today.isoformat(), "fetched_at": datetime.now(IL).isoformat(timespec="minutes"),
                "te": te or old.get("te") or {}, "curves": cv or old.get("curves") or {}}
        prices_changed = (snap["te"] != old.get("te")) or (snap["curves"] != old.get("curves"))
        snap_p.parent.mkdir(parents=True, exist_ok=True)
        snap_p.write_text(json.dumps(snap, ensure_ascii=False), encoding="utf-8")

    wb_p = OUT / "worldbank.json"
    try:
        wb_prev = json.loads(wb_p.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        wb_prev = {}
    wb, err = worldbank(cfg, creq.Session(impersonate="chrome"), wb_prev)
    if err:
        failures.append(f"הבנק העולמי: {err}")
    wb_changed = bool(wb) and wb is not wb_prev
    if wb_changed:
        wb_p.write_text(json.dumps(wb, ensure_ascii=False, indent=1), encoding="utf-8")

    new, feed_fail, meta = headlines(cfg, known, now)
    if new:
        write_items(new)
    if feed_fail and len(feed_fail) * 2 >= meta["sources"]:
        failures.append(f"{len(feed_fail)} מקורות כותרות מתוך {meta['sources']}")

    last = last_analysis()
    age_h = ((now - datetime.fromisoformat(last["analyzed_at"])).total_seconds() / 3600
             if last and last.get("analyzed_at") else 1e9)
    since = [r for r in meta["existing"] + new if not last or r["ts"] > (last.get("through") or "")]
    analyze = os.environ.get("COMMOD_FORCE") == "1" or (
        age_h >= MIN_HOURS and (len(since) >= MIN_NEW or prices_changed))

    (OUT / "state.json").write_text(json.dumps({
        "fetched_at": datetime.now(IL).isoformat(timespec="minutes"),
        "te_count": len(te), "curves": sorted(cv), "worldbank_updated": (wb or {}).get("updated"),
        "new_items": len(new), "failures": failures + feed_fail[:8]}, ensure_ascii=False, indent=1),
        encoding="utf-8")

    print(f"::notice::סחורות: {len(te)} מחירים מ-Trading Economics · עקומים {len(cv)} · "
          f"הבנק העולמי {(wb or {}).get('updated') or '—'} · {len(new)} כותרות חדשות · "
          f"ניתוח: {'כן' if analyze else 'לא'} ({len(since)} כותרות מאז הקודם"
          + (f", לפני {age_h:.1f} שעות" if age_h < 1e8 else ", אין ניתוח קודם") + ")")
    if failures:
        print("::warning title=מקורות סחורות שנכשלו::" + " · ".join(failures[:6]))
    print("  נפסלו: " + ", ".join(f"{k} {v}" for k, v in meta["dropped"].items()))

    gh = os.environ.get("GITHUB_OUTPUT")
    if gh:
        with open(gh, "a", encoding="utf-8") as f:
            f.write(f"new={len(new)}\n")
            f.write(f"analyze={'yes' if analyze else 'no'}\n")
            f.write(f"changed={'yes' if (new or prices_changed or wb_changed or analyze) else 'no'}\n")
    return 0 if (te or cv or new) else 1


if __name__ == "__main__":
    sys.exit(main())
